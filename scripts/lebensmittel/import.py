"""Importiert Lebensmittel in die Supabase-Tabelle `lebensmittel`.

Quellen:
  - Bundeslebensmittelschlüssel (BLS) 4.0, Max Rubner-Institut (2025), DOI 10.25826/Data20251217-134202-0
  - Open Food Facts (ODbL): in Deutschland meistgescannte Produkte

Läuft im GitHub-Workflow `lebensmittel-import.yml` und schreibt JSON-Teile nach `out/` (teil-01.json, …).
Der Workflow legt sie auf dem Branch `lebensmittel-daten` ab, Supabase übernimmt sie per pg_net
(siehe Migration 20261017150000_lebensmittel_import.sql).
Umgebung: optional OFF_ANZAHL (Standard 30000), BLS_URL.
"""

import csv
import gzip
import heapq
import io
import json
import math
import os
import re
import sys
import zipfile
from html import unescape

import openpyxl
import requests

UA = {'User-Agent': 'HLX-Coaching-App/1.0 (Lebensmittel-Import; github.com/YoungShindy/coaching-app)'}

BLS_GOVDATA = 'https://www.govdata.de/ckan/api/3/action/package_show?id=bundeslebensmittelschlussel-bls-version-4-0-deutsche-nahrstoffdatenbank'
BLS_SEITEN = ['https://www.openagrar.de/receive/openagrar_mods_00112643', 'https://blsdb.de/download']

# Hauptgruppen nach erstem Buchstaben des BLS-Codes
BLS_GRUPPEN = {
    'B': 'Brot & Backwaren', 'C': 'Getreide', 'D': 'Feine Backwaren', 'E': 'Eier', 'F': 'Obst',
    'G': 'Gemüse', 'H': 'Hülsenfrüchte & Nüsse', 'K': 'Kartoffeln & Pilze', 'M': 'Milch & Milchprodukte',
    'N': 'Getränke', 'P': 'Alkoholische Getränke', 'Q': 'Fette & Öle', 'R': 'Gewürze & Zutaten',
    'S': 'Süßwaren & Zucker', 'T': 'Fisch', 'U': 'Fleisch', 'V': 'Wild & Innereien', 'W': 'Wurst',
    'X': 'Gerichte', 'Y': 'Gerichte',
}

# Synonyme: steht der Schlüssel im BLS-Namen, sind auch diese Suchwörter treffend
SYNONYME = {
    'möhre': 'karotte mohrrübe', 'kartoffel': 'erdapfel', 'quark': 'topfen', 'brötchen': 'semmel schrippe',
    'pfannkuchen': 'eierkuchen crepe', 'tomate': 'paradeiser', 'blumenkohl': 'karfiol', 'aprikose': 'marille',
    'sahne': 'rahm', 'hackfleisch': 'hack faschiertes', 'pute': 'truthahn', 'hähnchen': 'huhn hühnchen chicken',
    'hühnerei': 'ei eier', 'rote bete': 'rote beete', 'feldsalat': 'rapunzel', 'teigwaren': 'nudeln pasta',
    'weizenmehl': 'mehl', 'haferflocken': 'porridge', 'pflaume': 'zwetschge', 'johannisbeere': 'ribisel',
    'meerrettich': 'kren', 'rosenkohl': 'kohlsprossen', 'grünkohl': 'braunkohl', 'kürbis': 'hokkaido',
}

PFLICHT = ['kcal', 'protein', 'kh', 'fett']


def zahl(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v) if math.isfinite(v) else None
    s = str(v).strip().replace(' ', '').replace(',', '.')
    if s in ('', '-', '—', 'n.a.', 'na', 'tr', 'Tr'):
        return None
    s = s.lstrip('<~≈ ')
    try:
        return float(s)
    except ValueError:
        return None


def runde(v, stellen=1):
    return None if v is None else round(v, stellen)


# ─── BLS ─────────────────────────────────────────────────────────────────────

SITZUNG = requests.Session()
SITZUNG.headers.update(UA)


def hrefs(html, basis):
    return [requests.compat.urljoin(basis, unescape(h)) for h in re.findall(r'href="([^"]+)"', html, re.I)]


def ist_datei_link(h):
    return bool(re.search(r'\.(xlsx|zip)(\?|$)', h, re.I)) or 'zipservlet' in h.lower() or 'download?token' in h.lower()


def bls_links():
    """Mögliche Download-Links der BLS-Datendatei sammeln (deutsche Datendatei zuerst)."""
    links = [os.environ.get('BLS_URL', '')]
    try:
        for res in SITZUNG.get(BLS_GOVDATA, timeout=60).json()['result']['resources']:
            links.append(res.get('url') or '')
    except (requests.RequestException, ValueError, KeyError) as e:
        print(f'govdata nicht nutzbar: {e}')
    for seite in BLS_SEITEN:
        try:
            gefunden = [h for h in hrefs(SITZUNG.get(seite, timeout=60).text, seite) if ist_datei_link(h)]
            print(f'{seite}: {gefunden[:10]}')
            links += gefunden
        except requests.RequestException as e:
            print(f'{seite} nicht erreichbar: {e}')
    links = [l for l in dict.fromkeys(links) if l]
    return sorted(links, key=lambda l: ('compon' in l.lower(), not re.search(r'daten|data', l, re.I), not re.search(r'_de\b|_de\.', l, re.I)))


def bls_aus_antwort(url, inhalt):
    """Excel-Datei aus einer Antwort holen (direkt oder aus ZIP); None, wenn keine passende Datei enthalten ist."""
    z = zipfile.ZipFile(io.BytesIO(inhalt))
    if 'xl/workbook.xml' in z.namelist():
        return inhalt if 'compon' not in url.lower() else None
    innen = sorted((n for n in z.namelist() if n.lower().endswith('.xlsx') and 'compon' not in n.lower()),
                   key=lambda n: (not re.search(r'daten|data', n, re.I), not re.search(r'_de\b|_de\.', n, re.I)))
    print(f'  ZIP-Inhalt: {z.namelist()[:10]}')
    return z.read(innen[0]) if innen else None


def bls_laden():
    besucht, offen = set(), bls_links()
    while offen:
        url = offen.pop(0)
        if url in besucht or len(besucht) > 30:
            continue
        besucht.add(url)
        try:
            r = SITZUNG.get(url, timeout=300)
        except requests.RequestException as e:
            print(f'BLS {url}: {e}')
            continue
        if r.ok and r.content[:2] == b'PK':
            datei = bls_aus_antwort(url, r.content)
            if datei:
                print(f'BLS geladen: {url} ({len(datei) // 1024} KB)')
                return datei
            continue
        art = r.headers.get('content-type', '')
        print(f'BLS {url}: HTTP {r.status_code}, {art}, Anfang: {r.text[:120]!r}')
        if r.ok and 'html' in art:
            # Verzeichnisseite: enthaltene Excel-Dateien vorne anstellen
            offen = [h for h in hrefs(r.text, url) if ist_datei_link(h) and h not in besucht] + offen
    sys.exit('BLS-Datei nicht gefunden. Link im Workflow-Feld „bls_url“ angeben.')


def spalte_finden(kopf, regeln, belegt):
    """Erste Spalte, deren Kopfzeile zu einer Regel passt (Regeln in Prioritätsreihenfolge)."""
    for regel in regeln:
        for i, k in enumerate(kopf):
            if i not in belegt and regel(k):
                return i
    return None


def wort(*w):
    return lambda k: any(re.search(rf'(^|[^a-z0-9]){re.escape(x)}([^a-z0-9]|$)', k) for x in w)


def enthaelt(alle, ohne=()):
    return lambda k: all(a in k for a in alle) and not any(o in k for o in ohne)


def bls_lesen(inhalt):
    wb = openpyxl.load_workbook(io.BytesIO(inhalt), read_only=True, data_only=True)
    blatt = max(wb.worksheets, key=lambda ws: ws.max_row or 0)
    zeilen = list(blatt.iter_rows(values_only=True))
    print(f'Tabellenblatt „{blatt.title}“: {len(zeilen)} Zeilen, {len(zeilen[0]) if zeilen else 0} Spalten')

    code_re = re.compile(r'^[A-Z]\d{6}$')

    def ist_code(z, j):
        return j < len(z) and isinstance(z[j], str) and bool(code_re.match(z[j].strip()))

    # Datenstart: erste Zeile, ab der fünf Zeilen in Folge in derselben Spalte einen BLS-Code haben
    treffer = next(((i, j) for i in range(min(len(zeilen) - 5, 50)) for j in range(len(zeilen[i]))
                    if all(ist_code(zeilen[i + k], j) for k in range(5))), None)
    if not treffer:
        for z in zeilen[:6]:
            print('  ', [str(c)[:20] for c in z[:10]])
        sys.exit('BLS-Codes nicht gefunden (erwartet: Buchstabe + 6 Ziffern).')
    start, code_sp = treffer
    breite = max(len(z) for z in zeilen[:start + 1])
    kopf = [' '.join(str(z[j]) for z in zeilen[:start] if j < len(z) and z[j] is not None).lower() for j in range(breite)]
    print(f'Datenstart in Zeile {start + 1}, BLS-Code in Spalte {code_sp + 1}')
    for j, k in enumerate(kopf[:12]):
        print(f'  Kopf {j + 1}: {k[:90]}')

    # Namensspalte: Text mit größter Durchschnittslänge in den ersten Datenzeilen
    probe = zeilen[start:start + 200]
    def textlaenge(j):
        werte = [z[j] for z in probe if j < len(z) and isinstance(z[j], str)]
        return sum(len(w) for w in werte) / max(1, len(probe)) if len(werte) > len(probe) * 0.8 else 0
    name_sp = next((j for j, k in enumerate(kopf) if 'bezeichnung' in k and 'referenz' not in k and 'herkunft' not in k), None)
    if name_sp is None:
        name_sp = max((j for j in range(breite) if j != code_sp), key=textlaenge)

    belegt = {code_sp, name_sp}
    regeln = {
        'kcal': [wort('enercc'), enthaelt(['energie', 'kcal']), enthaelt(['energy', 'kcal'])],
        'kj': [wort('enercj'), enthaelt(['energie', 'kj']), enthaelt(['energy', 'kj'])],
        'protein': [wort('prot625', 'prot'), enthaelt(['protein'], ['pflanz', 'tier', 'amino']), enthaelt(['eiweiß'], ['pflanz', 'tier'])],
        'fett': [wort('fat'), enthaelt(['fett'], ['säure', 'saeure', 'fettsäure'])],
        'kh': [wort('choavl', 'cho'), enthaelt(['kohlenhydrat'], ['ballast'])],
        'zucker': [wort('sugar'), enthaelt(['zucker', 'gesamt']), enthaelt(['zucker'], ['alkohol', 'austausch'])],
        'ballast': [wort('fibt'), enthaelt(['ballaststoff', 'gesamt']), enthaelt(['ballaststoff'])],
        'gesaettigt': [wort('fasat'), enthaelt(['gesättigt'], ['einfach', 'mehrfach', 'ungesättigt'])],
        'salz': [wort('nacl'), enthaelt(['salz'], ['salzig'])],
        'natrium': [wort('na'), enthaelt(['natrium'])],
    }
    sp = {}
    for feld, r in regeln.items():
        j = spalte_finden(kopf, r, belegt)
        if j is not None:
            sp[feld] = j
            belegt.add(j)
    print(f'Name in Spalte {name_sp + 1}: {kopf[name_sp][:60]}')
    for feld, j in sp.items():
        print(f'  {feld:<10} → Spalte {j + 1}: {kopf[j][:90]}')
    if 'kcal' not in sp and 'kj' in sp:
        print('  kcal wird aus kJ berechnet')
    fehlend = [f for f in PFLICHT if f not in sp and not (f == 'kcal' and 'kj' in sp)]
    if fehlend:
        sys.exit(f'Pflichtspalten nicht erkannt: {fehlend}')
    natrium_in_mg = 'natrium' in sp and 'mg' in kopf[sp['natrium']]

    def wert(z, feld):
        j = sp.get(feld)
        return zahl(z[j]) if j is not None and j < len(z) else None

    eintraege, gesehen = [], set()
    for z in zeilen[start:]:
        code = z[code_sp].strip() if code_sp < len(z) and isinstance(z[code_sp], str) else ''
        name = str(z[name_sp]).strip() if name_sp < len(z) and z[name_sp] else ''
        if not code_re.match(code) or not name or code in gesehen:
            continue
        gesehen.add(code)
        kcal = wert(z, 'kcal')
        if kcal is None and wert(z, 'kj') is not None:
            kcal = wert(z, 'kj') / 4.184
        if kcal is None:
            continue
        salz = wert(z, 'salz')
        if salz is None and wert(z, 'natrium') is not None:
            salz = wert(z, 'natrium') * 2.54 / (1000 if natrium_in_mg else 1)
        klein = name.lower()
        syn = ' '.join(v for k, v in SYNONYME.items() if k in klein) or None
        eintraege.append({
            'id': f'bls:{code}', 'quelle': 'bls', 'name': name[:200], 'marke': None,
            'kategorie': BLS_GRUPPEN.get(code[0]), 'barcode': None,
            'kcal': round(kcal), 'protein': runde(wert(z, 'protein')), 'kh': runde(wert(z, 'kh')),
            'zucker': runde(wert(z, 'zucker')), 'fett': runde(wert(z, 'fett')),
            'gesaettigt': runde(wert(z, 'gesaettigt')), 'ballast': runde(wert(z, 'ballast')),
            'salz': runde(salz, 2), 'ml': code[0] in 'NP', 'portion_g': None, 'menge': None,
            'nutriscore': None, 'nova': None, 'zusatzstoffe': None, 'bio': None,
            'beliebtheit': 0, 'suchwoerter': syn,
            'kategorien': None, 'obst_gemuese': 100 if code[0] in 'FG' else None, 'bild': None,
        })
    return eintraege


def bls_pruefen(eintraege):
    """Stichproben gegen bekannte Nährwerte, damit eine falsch erkannte Spalte nicht importiert wird."""
    proben = [
        ('banane', 'kcal', 70, 115), ('butter', 'fett', 70, 90), ('haferflocken', 'kh', 45, 72),
        ('vollmilch', 'protein', 2.8, 4.0), ('hühnerei', 'protein', 10, 15),
    ]
    fehler, gefunden = [], 0
    for wort_, feld, lo, hi in proben:
        treffer = sorted((e for e in eintraege if re.match(rf'{wort_}([ ,]|$)', e['name'].lower())), key=lambda e: len(e['name']))
        if not treffer:
            print(f'  Stichprobe {wort_}: nicht gefunden')
            continue
        e = treffer[0]
        gefunden += 1
        ok = e[feld] is not None and lo <= e[feld] <= hi
        print(f'  Stichprobe {e["name"]}: {feld}={e[feld]} kcal={e["kcal"]} P={e["protein"]} KH={e["kh"]} F={e["fett"]} {"✓" if ok else "✗"}')
        if not ok:
            fehler.append(e['name'])
    if fehler or gefunden < 3:
        sys.exit(f'Stichproben fehlgeschlagen (gefunden: {gefunden}, unplausibel: {fehler}). Spaltenzuordnung prüfen.')


# ─── Open Food Facts ─────────────────────────────────────────────────────────

OFF_EXPORT = 'https://static.openfoodfacts.org/data/en.openfoodfacts.org.products.csv.gz'
OFF_NAEHRWERTE = [
    'energy-kcal_100g', 'energy_100g', 'proteins_100g', 'carbohydrates_100g', 'fat_100g', 'sugars_100g',
    'saturated-fat_100g', 'fiber_100g', 'salt_100g', 'fruits-vegetables-legumes-estimate-from-ingredients_100g',
    'fruits-vegetables-nuts-estimate-from-ingredients_100g', 'fruits-vegetables-nuts_100g',
]


def liste(v):
    return [x for x in (v or '').split(',') if x]


def off_laden(ziel):
    """Die in Deutschland meistgescannten Produkte aus dem Komplett-Export (CSV, tab-getrennt, gzip)."""
    csv.field_size_limit(sys.maxsize)
    beste, zeilen, de, gesehen = [], 0, 0, set()
    with SITZUNG.get(OFF_EXPORT, stream=True, timeout=600) as r:
        r.raise_for_status()
        text = io.TextIOWrapper(gzip.GzipFile(fileobj=r.raw), encoding='utf-8', errors='replace', newline='')
        leser = csv.DictReader(text, delimiter='\t', quoting=csv.QUOTE_NONE)
        fehlend = [f for f in ['code', 'product_name', 'countries_tags', 'unique_scans_n', 'categories_tags'] + OFF_NAEHRWERTE[:4]
                   if f not in (leser.fieldnames or [])]
        print(f'OFF-Export: {len(leser.fieldnames or [])} Spalten, fehlend: {fehlend}')
        for z in leser:
            zeilen += 1
            if zeilen % 500000 == 0:
                print(f'  {zeilen} Zeilen gelesen, {de} aus Deutschland')
            if 'en:germany' not in (z.get('countries_tags') or ''):
                continue
            de += 1
            p = {
                'code': z.get('code'), 'product_name': z.get('product_name'), 'brands': z.get('brands'),
                'quantity': z.get('quantity'), 'serving_quantity': z.get('serving_quantity'),
                'nutriments': {k: z.get(k) for k in OFF_NAEHRWERTE},
                'nutriscore_grade': z.get('nutriscore_grade'), 'nova_group': z.get('nova_group'),
                'additives_tags': liste(z.get('additives_tags')), 'labels_tags': liste(z.get('labels_tags')),
                'unique_scans_n': z.get('unique_scans_n'), 'categories_tags': liste(z.get('categories_tags')),
                'image_front_small_url': z.get('image_front_small_url') or z.get('image_small_url'),
            }
            e = off_eintrag(p)
            if e and e['id'] not in gesehen:
                gesehen.add(e['id'])
                eintrag = (e['beliebtheit'], e['id'], e)
                if len(beste) < ziel:
                    heapq.heappush(beste, eintrag)
                elif eintrag[:2] > beste[0][:2]:
                    heapq.heapreplace(beste, eintrag)
    print(f'OFF-Export: {zeilen} Zeilen, {de} aus Deutschland, {len(beste)} übernommen')
    return [e for _, _, e in sorted(beste, key=lambda t: t[:2], reverse=True)]


def off_eintrag(p):
    code = str(p.get('code') or '').strip()
    name = (p.get('product_name_de') or p.get('product_name') or '').strip()
    n = p.get('nutriments') or {}
    if not re.fullmatch(r'\d{8,14}', code) or not name:
        return None
    kcal = zahl(n.get('energy-kcal_100g'))
    if kcal is None and zahl(n.get('energy_100g')) is not None:
        kcal = zahl(n.get('energy_100g')) / 4.184
    protein, kh, fett = zahl(n.get('proteins_100g')), zahl(n.get('carbohydrates_100g')), zahl(n.get('fat_100g'))
    if kcal is None or kcal < 0 or kcal > 900 or None in (protein, kh, fett):
        return None
    if protein + kh + fett > 105 or min(protein, kh, fett) < 0:
        return None
    menge = (p.get('quantity') or '').strip()[:40] or None
    ml = bool(menge and re.search(r'\d\s*(ml|cl|l)\b', menge.lower()))
    labels = p.get('labels_tags') or []
    nova = p.get('nova_group')
    grade = (p.get('nutriscore_grade') or '').lower()
    return {
        'id': f'off:{code}', 'quelle': 'off', 'name': name[:200], 'marke': (p.get('brands') or '').split(',')[0].strip()[:80] or None,
        'kategorie': None, 'barcode': code,
        'kcal': round(kcal), 'protein': runde(protein), 'kh': runde(kh), 'zucker': runde(zahl(n.get('sugars_100g'))),
        'fett': runde(fett), 'gesaettigt': runde(zahl(n.get('saturated-fat_100g'))), 'ballast': runde(zahl(n.get('fiber_100g'))),
        'salz': runde(zahl(n.get('salt_100g')), 2), 'ml': ml,
        'portion_g': runde(zahl(p.get('serving_quantity'))) if (zahl(p.get('serving_quantity')) or 0) > 0 else None,
        'menge': menge, 'nutriscore': grade if grade in ('a', 'b', 'c', 'd', 'e') else None,
        'nova': int(zahl(nova)) if zahl(nova) in (1, 2, 3, 4) else None,
        'zusatzstoffe': [a.split(':', 1)[-1] for a in (p.get('additives_tags') or [])][:40] or None,
        'bio': any(l in ('en:organic', 'en:eu-organic', 'de:bio') for l in labels),
        'beliebtheit': int(zahl(p.get('unique_scans_n')) or 0), 'suchwoerter': None,
        # Für Produkt-Check und Alternativen: die drei genauesten Kategorien, Obst-/Gemüseanteil, Bild
        'kategorien': (p.get('categories_tags') or [])[-3:] or None,
        'obst_gemuese': runde(next((zahl(n.get(k)) for k in (
            'fruits-vegetables-legumes-estimate-from-ingredients_100g', 'fruits-vegetables-nuts-estimate-from-ingredients_100g',
            'fruits-vegetables-nuts_100g') if zahl(n.get(k)) is not None), None)),
        'bild': (p.get('image_front_small_url') or '')[:300] or None,
    }


# ─── Ausgabe ─────────────────────────────────────────────────────────────────

def schreiben(eintraege, ordner='out', je_teil=3000):
    os.makedirs(ordner, exist_ok=True)
    teile = 0
    for i in range(0, len(eintraege), je_teil):
        teile += 1
        with open(f'{ordner}/teil-{teile:02d}.json', 'w', encoding='utf-8') as f:
            json.dump(eintraege[i:i + je_teil], f, ensure_ascii=False, separators=(',', ':'))
    with open(f'{ordner}/teile.txt', 'w') as f:
        f.write(str(teile))
    print(f'{len(eintraege)} Einträge in {teile} Teilen geschrieben')


def main():
    bls = bls_lesen(bls_laden())
    print(f'BLS: {len(bls)} Lebensmittel erkannt')
    bls_pruefen(bls)
    if len(bls) < 5000:
        sys.exit('Zu wenige BLS-Einträge, Import abgebrochen.')

    off = off_laden(int(os.environ.get('OFF_ANZAHL') or 30000))
    print(f'Open Food Facts: {len(off)} Produkte')
    for e in off[:5]:
        print(f'  {e["name"]} ({e["marke"]}): {e["kcal"]} kcal, Scans {e["beliebtheit"]}')
    schreiben(bls + off)


if __name__ == '__main__':
    main()
